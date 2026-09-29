"""Tests for the branded DOCX renderer and the PDF conversion."""

from __future__ import annotations

import copy
import logging
import os
import re
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import Cm, Emu

from wiseglow_desk.docgen.brand import Brand
from wiseglow_desk.docgen.docx_renderer import add_rich_text, format_date, render_docx
from wiseglow_desk.docgen.model import Document
from wiseglow_desk.docgen.pdf import docx_to_pdf, find_soffice, pdf_page_count

MINIMAL: dict[str, Any] = {
    "document_id": "min-doc",
    "doc_type": "GENERIC",
    "title": "Documento minimo",
    "sections": [{"heading": "Sezione", "blocks": [{"type": "paragraph", "text": "Testo."}]}],
}


def _xml(path: Path, member: str) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read(member).decode("utf-8")


def _all_text(path: Path) -> str:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


@pytest.mark.parametrize("name", ["nbo", "loi", "interesse_terreno"])
def test_render_samples(sample_dicts: dict[str, Any], brand: Brand, tmp_path: Path, name: str) -> None:
    doc = Document.from_dict(sample_dicts[name])
    out = render_docx(doc, brand, tmp_path / f"{name}.docx")
    assert out.is_file() and out.stat().st_size > 10_000
    text = _all_text(out)
    assert doc.title in text
    for section in doc.sections:
        assert section.heading in text
    for sig in doc.signatures:
        assert sig.party in text
    for att in doc.attachments:
        assert att in text
    core = DocxDocument(str(out)).core_properties
    assert core.title == doc.title
    assert core.author == brand.company.name
    assert core.subject == doc.doc_type
    assert core.keywords == doc.document_id


def test_header_contains_logo_and_rule(sample_dicts: dict[str, Any], brand: Brand, tmp_path: Path) -> None:
    out = render_docx(Document.from_dict(sample_dicts["nbo"]), brand, tmp_path / "h.docx")
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
    assert any(n.startswith("word/media/") for n in names)
    header = _xml(out, "word/header1.xml")
    assert "<w:drawing>" in header
    assert brand.company.name in header
    assert brand.colors.primary.lstrip("#") in header  # rule colour
    assert brand.footer.confidential_label.upper() in header  # nbo is confidential


def test_footer_has_page_fields_and_labels(sample_dicts: dict[str, Any], brand: Brand, tmp_path: Path) -> None:
    out = render_docx(Document.from_dict(sample_dicts["nbo"]), brand, tmp_path / "f.docx")
    footer = _xml(out, "word/footer1.xml")
    assert " PAGE " in footer and " NUMPAGES " in footer
    assert brand.footer.text in footer
    assert "Pagina" in footer and brand.footer.confidential_label in footer


def test_missing_logo_falls_back_to_company_name(brand: Brand, tmp_path: Path) -> None:
    b = Brand.from_dict(brand.to_dict(), base_dir=tmp_path)  # logo.png does not exist here
    assert not b.logo_path.exists()
    out = render_docx(Document.from_dict(MINIMAL), b, tmp_path / "nologo.docx")
    header = _xml(out, "word/header1.xml")
    assert "<w:drawing>" not in header
    assert header.count(b.company.name) >= 2


def test_minimal_document_renders(brand: Brand, tmp_path: Path) -> None:
    out = render_docx(Document.from_dict(MINIMAL), brand, tmp_path / "min.docx")
    text = _all_text(out)
    assert "Documento minimo" in text and "Sezione" in text and "Testo." in text
    assert "Firme" not in text and "Allegati" not in text


def test_all_block_types_render(brand: Brand, tmp_path: Path) -> None:
    data = copy.deepcopy(MINIMAL)
    data["sections"][0]["blocks"] = [
        {"type": "paragraph", "text": "Para **bold** e *corsivo*.\nSeconda riga."},
        {"type": "bullets", "items": ["uno", "due"]},
        {"type": "numbered", "items": ["primo", "secondo"]},
        {"type": "table", "columns": ["A", "B", "C"], "rows": [["1", "2", "3"], ["4", "5"]], "col_widths": [20, 30, 50], "caption": "Didascalia"},
        {"type": "callout", "text": "Attenzione", "style": "warning"},
        {"type": "callout", "text": "Nota", "style": "note"},
        {"type": "quote", "text": "Citazione"},
        {"type": "page_break"},
        {"type": "paragraph", "text": "Dopo il salto."},
    ]
    data["status"] = "Finale"
    out = render_docx(Document.from_dict(data), brand, tmp_path / "blocks.docx")
    text = _all_text(out)
    for needle in ("Para bold e corsivo.", "uno", "secondo", "Didascalia", "Attenzione", "Nota", "Citazione", "Dopo il salto.", "FINALE"):
        assert needle in text
    body = _xml(out, "word/document.xml")
    assert 'w:type="page"' in body
    assert brand.colors.primary.lstrip("#") in body  # table header fill
    assert brand.colors.table_alt_row.lstrip("#") in body


def test_table_col_widths_are_proportional(brand: Brand, tmp_path: Path) -> None:
    data = copy.deepcopy(MINIMAL)
    data["sections"][0]["blocks"] = [{"type": "table", "columns": ["A", "B"], "rows": [["1", "2"]], "col_widths": [25, 75]}]
    out = render_docx(Document.from_dict(data), brand, tmp_path / "w.docx")
    table = DocxDocument(str(out)).tables[-1]
    w0, w1 = (c.width for c in table.rows[0].cells)
    assert abs(w1 / w0 - 3.0) < 0.05
    m = brand.page.margins_cm
    assert abs(Emu(w0 + w1).cm - (21.0 - m.left - m.right)) < 0.05


def test_numbered_lists_restart(brand: Brand, tmp_path: Path) -> None:
    data = copy.deepcopy(MINIMAL)
    data["sections"][0]["blocks"] = [
        {"type": "numbered", "items": ["a", "b"]},
        {"type": "paragraph", "text": "x"},
        {"type": "numbered", "items": ["c"]},
    ]
    out = render_docx(Document.from_dict(data), brand, tmp_path / "num.docx")
    doc = DocxDocument(str(out))
    num_ids = []
    for p in doc.paragraphs:
        if p.style.name == "List Number" and p._p.pPr is not None and p._p.pPr.numPr is not None:
            num_ids.append(p._p.pPr.numPr.numId.val)
    assert len(num_ids) == 3 and num_ids[0] == num_ids[1] != num_ids[2]


def test_deterministic_output(sample_dicts: dict[str, Any], brand: Brand, tmp_path: Path) -> None:
    doc = Document.from_dict(sample_dicts["loi"])
    a = render_docx(doc, brand, tmp_path / "a.docx").read_bytes()
    b = render_docx(doc, brand, tmp_path / "b.docx").read_bytes()
    assert a == b


def test_east_asian_font_and_chinese_labels(brand: Brand, tmp_path: Path) -> None:
    data = copy.deepcopy(MINIMAL)
    data.update({"language": "zh", "title": "非约束性报价", "reference": "WG-001", "confidential": True})
    data["sections"] = [{"heading": "一、前言", "blocks": [{"type": "paragraph", "text": "我们是 **WiseGlow**。"}]}]
    data["signatures"] = [{"party": "WiseGlow S.r.l.", "name": "张伟", "title": "总经理"}]
    out = render_docx(Document.from_dict(data), brand, tmp_path / "zh.docx")
    doc = DocxDocument(str(out))
    rfonts = doc.styles["Normal"].element.rPr.find(qn("w:rFonts"))
    assert rfonts.get(qn("w:eastAsia")) == brand.fonts.body_cjk
    assert rfonts.get(qn("w:ascii")) == brand.fonts.body
    h1 = doc.styles["Heading 1"].element.rPr.find(qn("w:rFonts"))
    assert h1.get(qn("w:ascii")) == brand.fonts.heading and h1.get(qn("w:asciiTheme")) is None
    footer = _xml(out, "word/footer1.xml")
    assert "第" in footer and "页，共" in footer
    text = _all_text(out)
    assert "签署" in text and "编号" in text and "日期" in text
    assert "年" in text and "月" in text and "日" in text  # localised date


@pytest.mark.parametrize("name", ["nbo", "loi", "interesse_terreno"])
def test_property_children_follow_ooxml_order(sample_dicts: dict[str, Any], brand: Brand, tmp_path: Path, name: str) -> None:
    """Word (unlike LibreOffice) rejects property elements whose children are out of order."""
    from xml.etree import ElementTree as ET

    from wiseglow_desk.docgen.docx_renderer import _SCHEMA_ORDER

    out = render_docx(Document.from_dict(sample_dicts[name]), brand, tmp_path / f"{name}.docx")
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(out) as z:
        parts = [n for n in z.namelist() if n.startswith("word/") and n.endswith(".xml")]
        for part in parts:
            root = ET.fromstring(z.read(part))
            for parent in root.iter():
                seq = _SCHEMA_ORDER.get("w:" + parent.tag.replace(ns, ""))
                if not seq:
                    continue
                positions = [seq.index("w:" + c.tag.replace(ns, "")) for c in parent if c.tag.startswith(ns) and "w:" + c.tag.replace(ns, "") in seq]
                assert positions == sorted(positions), f"{part}: {parent.tag} children out of order: {[c.tag for c in parent]}"


def test_add_rich_text_markdown(tmp_path: Path) -> None:
    doc = DocxDocument()
    p = doc.add_paragraph()
    runs = add_rich_text(p, "a **b** *c*\nd", size=10)
    texts = [(r.text, bool(r.bold), bool(r.italic)) for r in runs]
    # the run carrying the <w:br/> reads back as "\n"
    assert texts == [("a ", False, False), ("b", True, False), (" ", False, False), ("c", False, True), ("\n", False, False), ("d", False, False)]
    assert p.text == "a b c\nd"
    p2 = doc.add_paragraph()
    add_rich_text(p2, "2 * 3 * 4 e **grassetto con *stella*** fine")
    assert p2.text == "2 * 3 * 4 e grassetto con *stella* fine"


def test_format_date() -> None:
    assert format_date("2026-09-29", "it") == "29 settembre 2026"
    assert format_date("2026-09-29", "en") == "29 September 2026"
    assert format_date("2026-09-29", "zh") == "2026年9月29日"
    assert format_date("boh", "it") == "boh"


# -- PDF ---------------------------------------------------------------------


def test_find_soffice_env_override(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    fake = tmp_path / "soffice"
    fake.write_text("#!/bin/sh\n")
    monkeypatch.setenv("WISEGLOW_SOFFICE", str(fake))
    assert find_soffice() == str(fake)


def test_docx_to_pdf_without_soffice(monkeypatch: pytest.MonkeyPatch, brand: Brand, tmp_path: Path) -> None:
    monkeypatch.setattr("wiseglow_desk.docgen.pdf.find_soffice", lambda: None)
    out = render_docx(Document.from_dict(MINIMAL), brand, tmp_path / "x.docx")
    assert docx_to_pdf(out, tmp_path) is None


@pytest.mark.parametrize("name", ["nbo", "loi", "interesse_terreno"])
def test_pdf_render_samples(require_pdf: None, sample_dicts: dict[str, Any], brand: Brand, tmp_path: Path, name: str) -> None:
    doc = Document.from_dict(sample_dicts[name])
    docx_path = render_docx(doc, brand, tmp_path / f"{name}.docx")
    pdf = docx_to_pdf(docx_path, tmp_path, timeout=180)
    assert pdf is not None and pdf.is_file()
    pages = pdf_page_count(pdf)
    assert 2 <= pages <= 5, f"{name}: {pages} pagine"


def test_list_levels_share_the_style_indents(brand: Brand, tmp_path: Path) -> None:
    """Marker, first line and wrapped lines line up in Word and LibreOffice alike.

    Word applies the numbering level's indent/tab over the paragraph style,
    LibreOffice only when numPr sits on the paragraph: both must agree.
    """
    data = copy.deepcopy(MINIMAL)
    data["sections"][0]["blocks"] = [
        {"type": "bullets", "items": ["uno", "due"]},
        {"type": "numbered", "items": ["tre"]},
    ]
    out = render_docx(Document.from_dict(data), brand, tmp_path / "lists.docx")
    doc = DocxDocument(str(out))
    numbering = doc.part.numbering_part.element
    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        left = style.paragraph_format.left_indent
        hanging = Emu(-style.paragraph_format.first_line_indent)
        assert left.twips > 0 and 0 < hanging.twips < left.twips
        num = numbering.num_having_numId(style.element.pPr.numPr.numId.val)
        abstract = next(
            a for a in numbering.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == str(num.abstractNumId.val)
        )
        lvl0 = next(l for l in abstract.findall(qn("w:lvl")) if l.get(qn("w:ilvl")) == "0")
        ppr = lvl0.find(qn("w:pPr"))
        ind = ppr.find(qn("w:ind"))
        assert int(ind.get(qn("w:left"))) == left.twips
        assert int(ind.get(qn("w:hanging"))) == hanging.twips
        tabs = [int(t.get(qn("w:pos"))) for t in ppr.iter(qn("w:tab"))]
        assert tabs and all(pos == left.twips for pos in tabs)


# -- review fixes ------------------------------------------------------------


def test_add_rich_text_bold_italic() -> None:
    doc = DocxDocument()
    p = doc.add_paragraph()
    runs = add_rich_text(p, "Prezzo: ***EUR 1.000***.")
    assert [(r.text, bool(r.bold), bool(r.italic)) for r in runs] == [
        ("Prezzo: ", False, False),
        ("EUR 1.000", True, True),
        (".", False, False),
    ]
    assert "*" not in p.text


def _header_extent(path: Path) -> tuple[int, int]:
    m = re.search(r'<wp:extent cx="(\d+)" cy="(\d+)"/>', _xml(path, "word/header1.xml"))
    assert m, "nessuna immagine nell'intestazione"
    return int(m.group(1)), int(m.group(2))


@pytest.mark.parametrize("size", [(600, 600), (400, 800)])
def test_header_logo_height_is_capped(tmp_path: Path, size: tuple[int, int]) -> None:
    """A square/portrait real logo must not grow the header (repeated on every page)."""
    from PIL import Image

    Image.new("RGB", size, "white").save(tmp_path / "logo.png")
    brand = Brand.from_dict({}, base_dir=tmp_path)  # defaults: width 4.2 cm, max height 1.6 cm
    out = render_docx(Document.from_dict(MINIMAL), brand, tmp_path / "logo.docx")
    cx, cy = _header_extent(out)
    assert cy <= Cm(brand.logo.max_height_cm) and cx < Cm(brand.logo.width_cm)
    assert abs(cx / cy - size[0] / size[1]) < 0.01  # aspect ratio kept


def test_header_wide_logo_keeps_brand_width(brand: Brand, tmp_path: Path) -> None:
    out = render_docx(Document.from_dict(MINIMAL), brand, tmp_path / "wide.docx")
    cx, cy = _header_extent(out)
    assert cx == Cm(brand.logo.width_cm) and cy <= Cm(brand.logo.max_height_cm)


def test_borders_follow_schema_order_without_duplicates(brand: Brand, tmp_path: Path) -> None:
    blocks = [
        {"type": "callout", "text": "Nota", "style": "info"},
        {"type": "table", "columns": ["a", "b"], "rows": [["1", "2"]]},
    ]
    doc = dict(MINIMAL, sections=[{"heading": "S", "blocks": blocks}])
    out = render_docx(Document.from_dict(doc), brand, tmp_path / "b.docx")
    body = _xml(out, "word/document.xml")
    cell_borders = re.findall(r"<w:tcBorders>(.*?)</w:tcBorders>", body, re.S)
    assert cell_borders and all(re.findall(r"<w:(\w+)\b", b) == ["top", "left", "bottom", "right"] for b in cell_borders)
    for tbl in re.findall(r"<w:tblBorders>(.*?)</w:tblBorders>", body, re.S):
        sides = re.findall(r"<w:(\w+)\b", tbl)
        assert sides == ["top", "left", "bottom", "right", "insideH", "insideV"]


def test_find_soffice_warns_on_invalid_env(monkeypatch: pytest.MonkeyPatch, tmp_path: Path, caplog: pytest.LogCaptureFixture) -> None:
    bogus = str(tmp_path / "LibreOffice.app")
    monkeypatch.setenv("WISEGLOW_SOFFICE", bogus)
    with caplog.at_level(logging.WARNING, logger="wiseglow_desk.docgen.pdf"):
        assert find_soffice() != bogus
    assert bogus in caplog.text and "lo ignoro" in caplog.text


def _alive(pid: int) -> bool:
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    try:
        state = Path(f"/proc/{pid}/stat").read_text().rsplit(")", 1)[-1].split()[0]
    except OSError:  # no /proc (macOS) or the process vanished meanwhile
        return not os.path.isdir("/proc")
    return state != "Z"  # a zombie is dead, just not reaped yet


@pytest.mark.skipif(sys.platform == "win32", reason="process groups are POSIX")
def test_docx_to_pdf_timeout_kills_grandchildren(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """soffice is a wrapper (→ oosplash → soffice.bin): a timeout must not leave orphans or the profile dir."""
    pidfile = tmp_path / "child.pid"
    fake = tmp_path / "soffice"
    fake.write_text(f"#!/bin/sh\nsleep 60 &\necho $! > '{pidfile}'\nwait\n", encoding="utf-8")
    fake.chmod(0o755)
    monkeypatch.setenv("WISEGLOW_SOFFICE", str(fake))
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))  # profile dirs land here
    docx = tmp_path / "x.docx"
    docx.write_bytes(b"not really a docx")

    assert docx_to_pdf(docx, tmp_path / "out", timeout=1) is None
    pid = int(pidfile.read_text().strip())
    deadline = time.monotonic() + 5
    while _alive(pid) and time.monotonic() < deadline:
        time.sleep(0.05)
    assert not _alive(pid), "il processo nipote è sopravvissuto al timeout"
    assert not list(tmp_path.glob("wiseglow-soffice-*"))


def test_docx_to_pdf_reports_failed_conversion(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    fake = tmp_path / "soffice"
    fake.write_text("#!/bin/sh\necho 'boom' >&2\nexit 3\n", encoding="utf-8")
    fake.chmod(0o755)
    monkeypatch.setenv("WISEGLOW_SOFFICE", str(fake))
    docx = tmp_path / "x.docx"
    docx.write_bytes(b"x")
    assert docx_to_pdf(docx, tmp_path / "out", timeout=10) is None
