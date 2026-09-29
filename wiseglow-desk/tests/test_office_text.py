"""Tests for office_text.extract_text."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from docx import Document as DocxDocument

from wiseglow_desk.docgen.office_text import TRUNCATION_MARKER, UnsupportedFileError, extract_text


def test_docx_paragraphs_and_tables_in_order(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("Primo paragrafo")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Voce", "Valore"
    t.cell(1, 0).text, t.cell(1, 1).text = "Prezzo", "€ 100"
    doc.add_paragraph("Dopo la tabella")
    path = tmp_path / "x.docx"
    doc.save(str(path))
    text = extract_text(path)
    assert "Primo paragrafo" in text and "| Voce | Valore |" in text and "| Prezzo | € 100 |" in text
    assert text.index("Primo paragrafo") < text.index("| Voce") < text.index("Dopo la tabella")


def test_xlsx_sheets_as_tsv(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dati"
    ws.append(["Nome", "kWp", "Prezzo"])
    ws.append(["Impianto A", 120, 95000.5])
    ws.append([None, None, None])
    ws2 = wb.create_sheet("Note")
    ws2["A1"] = "solo una cella"
    path = tmp_path / "x.xlsx"
    wb.save(str(path))
    text = extract_text(path)
    assert "## Foglio: Dati" in text and "## Foglio: Note" in text
    assert "Nome\tkWp\tPrezzo" in text and "Impianto A\t120\t95000.5" in text
    assert "solo una cella" in text
    assert text.count("\n\n") == 1  # empty row skipped, blocks separated once


@pytest.mark.parametrize("suffix", [".csv", ".txt", ".md"])
def test_plain_files(tmp_path: Path, suffix: str) -> None:
    path = tmp_path / f"f{suffix}"
    path.write_bytes("﻿a;b\nciao;mondo\n".encode("utf-8"))
    assert extract_text(path) == "a;b\nciao;mondo"


def test_unsupported_and_missing(tmp_path: Path) -> None:
    for name in ("slide.pptx", "doc.pdf", "img.png"):
        p = tmp_path / name
        p.write_bytes(b"x")
        with pytest.raises(UnsupportedFileError, match="non supportato"):
            extract_text(p)
    with pytest.raises(FileNotFoundError):
        extract_text(tmp_path / "manca.txt")


def test_truncation_marker(tmp_path: Path) -> None:
    path = tmp_path / "long.txt"
    path.write_text("parola " * 1000, encoding="utf-8")
    text = extract_text(path, max_chars=200)
    assert len(text) <= 200
    assert text.endswith("\n" + TRUNCATION_MARKER)
    assert extract_text(path, max_chars=100_000) == ("parola " * 1000).strip()


def test_cp1252_csv_keeps_euro_and_accents(tmp_path: Path) -> None:
    """Excel (Windows, Italian locale) saves plain CSV in cp1252, not UTF-8."""
    path = tmp_path / "listino.csv"
    path.write_bytes("Caparra;€ 66.550,00\nCittà;Nardò\n".encode("cp1252"))
    assert extract_text(path) == "Caparra;€ 66.550,00\nCittà;Nardò"
    utf8 = tmp_path / "u.csv"
    utf8.write_text("Città;Nardò\n", encoding="utf-8")
    assert extract_text(utf8) == "Città;Nardò"
