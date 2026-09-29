"""DOCX renderer: the single place where the WiseGlow visual identity lives.

``render_docx(document, brand, out_path)`` turns a validated :class:`Document`
into a branded Word file with python-docx.  Everything visual (fonts, colours,
header/footer, tables, callouts, signatures) is decided here from the
:class:`Brand`; the document only carries content.

Output is deterministic for the same input: the zip is re-packed with fixed
entry timestamps and core properties use the document date, never "now".
"""

from __future__ import annotations

import datetime as dt
import io
import logging
import re
import zipfile
from pathlib import Path
from typing import Any, Iterable

from docx import Document as DocxDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Length, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .brand import Brand
from .catalog import DOC_TYPES
from .model import Block, Document, Section

log = logging.getLogger(__name__)

_PAGE_SIZES_CM: dict[str, tuple[float, float]] = {"A4": (21.0, 29.7), "LETTER": (21.59, 27.94)}
_MONTHS_IT = [
    "gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno",
    "luglio", "agosto", "settembre", "ottobre", "novembre", "dicembre",
]
_MONTHS_EN = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]
# ***bold italic***, **bold**, *italic* (a lone "*" between words is left alone)
_MD_TOKEN = re.compile(r"\*\*\*(.+?)\*\*\*|\*\*(.+?)\*\*|(?<![\w*])\*(?!\s)([^*\n]+?)\*(?![\w*])", re.S)
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
# List items: text (and wrapped lines) start at LEFT; the bullet/number sits HANGING before it.
_LIST_LEFT_CM = 1.0
_LIST_HANGING_CM = 0.6

# OOXML child order of the property elements we write into. Word rejects files
# whose children are out of sequence, LibreOffice does not: always insert in order.
_SCHEMA_ORDER: dict[str, tuple[str, ...]] = {
    "w:pPr": (
        "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl", "w:numPr",
        "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
        "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
        "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
        "w:sectPr", "w:pPrChange",
    ),
    "w:rPr": (
        "w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike", "w:dstrike",
        "w:outline", "w:shadow", "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden",
        "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect",
        "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
        "w:specVanish", "w:oMath",
    ),
    "w:tblPr": (
        "w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize", "w:tblStyleColBandSize",
        "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar",
        "w:tblLook", "w:tblCaption", "w:tblDescription",
    ),
    "w:tcPr": (
        "w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd", "w:noWrap", "w:tcMar",
        "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
    ),
    "w:trPr": (
        "w:cnfStyle", "w:divId", "w:gridBefore", "w:gridAfter", "w:wBefore", "w:wAfter", "w:cantSplit", "w:trHeight",
        "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden",
    ),
    "w:pBdr": ("w:top", "w:left", "w:bottom", "w:right", "w:between", "w:bar"),
    "w:tblBorders": ("w:top", "w:start", "w:left", "w:bottom", "w:end", "w:right", "w:insideH", "w:insideV"),
    "w:tcBorders": (
        "w:top", "w:start", "w:left", "w:bottom", "w:end", "w:right", "w:insideH", "w:insideV", "w:tl2br", "w:tr2bl",
    ),
    "w:tblCellMar": ("w:top", "w:start", "w:left", "w:bottom", "w:end", "w:right"),
}


def _local_tag(element: Any) -> str:
    return "w:" + str(element.tag).split("}")[-1]


def _insert_ordered(parent: Any, child: Any) -> Any:
    """Insert ``child`` into ``parent`` at its schema position (append if unknown)."""
    seq = _SCHEMA_ORDER.get(_local_tag(parent), ())
    tag = _local_tag(child)
    if tag in seq:
        successors = {qn(t) for t in seq[seq.index(tag) + 1 :]}
        for existing in parent:
            if existing.tag in successors:
                existing.addprevious(child)
                return child
    parent.append(child)
    return child


# ----------------------------------------------------------------------------
# public API
# ----------------------------------------------------------------------------


def render_docx(document: Document, brand: Brand, out_path: Path | str) -> Path:
    """Render ``document`` with ``brand`` into ``out_path`` (.docx) and return it."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    renderer = _Renderer(document, brand)
    renderer.build()
    renderer.save(out_path)
    return out_path


def add_rich_text(paragraph: Paragraph, text: str, **style: Any) -> list[Run]:
    """Add ``text`` to ``paragraph`` interpreting ``**bold**``, ``*italic*`` and ``\\n``.

    ``style`` keywords (font, size, color, bold, italic, cjk) are applied to
    every run created; markdown bold/italic are OR-ed on top.  Returns the runs.
    """
    runs: list[Run] = []
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    pos = 0
    for m in _MD_TOKEN.finditer(text):
        if m.start() > pos:
            runs += _add_plain(paragraph, text[pos : m.start()], style)
        if m.group(1) is not None:
            runs += _add_plain(paragraph, m.group(1), {**style, "bold": True, "italic": True})
        elif m.group(2) is not None:
            runs += _add_plain(paragraph, m.group(2), {**style, "bold": True})
        else:
            runs += _add_plain(paragraph, m.group(3), {**style, "italic": True})
        pos = m.end()
    if pos < len(text):
        runs += _add_plain(paragraph, text[pos:], style)
    return runs


def format_date(iso: str, language: str) -> str:
    """Localised long date: ``29 settembre 2026`` / ``29 September 2026`` / ``2026年9月29日``."""
    try:
        d = dt.date.fromisoformat(iso)
    except (TypeError, ValueError):
        return iso or ""
    if language == "zh":
        return f"{d.year}年{d.month}月{d.day}日"
    if language == "en":
        return f"{d.day} {_MONTHS_EN[d.month - 1]} {d.year}"
    return f"{d.day} {_MONTHS_IT[d.month - 1]} {d.year}"


# ----------------------------------------------------------------------------
# low-level helpers (XML)
# ----------------------------------------------------------------------------


def _add_plain(paragraph: Paragraph, text: str, style: dict[str, Any]) -> list[Run]:
    """Add plain text, turning newlines into line breaks; one run per line."""
    runs: list[Run] = []
    lines = text.split("\n")
    for i, line in enumerate(lines):
        run = paragraph.add_run(line)
        _style_run(run, **style)
        if i < len(lines) - 1:
            run.add_break()
        runs.append(run)
    return runs


def _style_run(
    run: Run,
    *,
    font: str | None = None,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    cjk: str | None = None,
) -> Run:
    if font or cjk:
        _set_fonts(run._r, font, cjk)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = _rgb(color)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    return run


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _blend(hex_color: str, ratio: float, base: str = "#FFFFFF") -> str:
    """Mix ``hex_color`` into ``base`` (``ratio`` = share of hex_color)."""
    c = [int(hex_color.lstrip("#")[i : i + 2], 16) for i in (0, 2, 4)]
    b = [int(base.lstrip("#")[i : i + 2], 16) for i in (0, 2, 4)]
    mixed = [round(b[i] + (c[i] - b[i]) * ratio) for i in range(3)]
    return "#{:02X}{:02X}{:02X}".format(*mixed)


def _set_fonts(element: Any, latin: str | None, cjk: str | None) -> None:
    """Set w:rFonts on a run/style element (ascii, hAnsi, cs, eastAsia); drop theme fonts."""
    rpr = element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = _insert_ordered(rpr, OxmlElement("w:rFonts"))
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    if latin:
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), latin)
    if cjk:
        rfonts.set(qn("w:eastAsia"), cjk)


def _set_lang(element: Any, latin: str, cjk: str = "zh-CN") -> None:
    rpr = element.get_or_add_rPr()
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = _insert_ordered(rpr, OxmlElement("w:lang"))
    lang.set(qn("w:val"), latin)
    lang.set(qn("w:eastAsia"), cjk)


def _para_border(paragraph: Paragraph, side: str, size_pt: float, color: str, space: int = 1) -> None:
    """Add a single paragraph border (side: top|bottom|left|right)."""
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = _insert_ordered(ppr, OxmlElement("w:pBdr"))
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(int(size_pt * 8)))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color.lstrip("#").upper())
    _insert_ordered(pbdr, el)


def _shade(cell: _Cell, fill: str) -> None:
    _insert_ordered(cell._tc.get_or_add_tcPr(), _shd_element(fill))


def _run_shade(run: Run, fill: str) -> None:
    _insert_ordered(run._r.get_or_add_rPr(), _shd_element(fill))


def _shd_element(fill: str) -> Any:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill.lstrip("#").upper())
    return shd


def _border_el(side: str, size_pt: float | None, color: str) -> Any:
    el = OxmlElement(f"w:{side}")
    if size_pt is None:
        el.set(qn("w:val"), "nil")
        return el
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(int(size_pt * 8)))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color.lstrip("#").upper())
    return el


def _table_borders(table: Table, size_pt: float | None, color: str, inside: bool = True) -> None:
    """Set all table borders (``size_pt=None`` removes them); earlier settings are replaced."""
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = _insert_ordered(tblpr, OxmlElement("w:tblBorders"))
    for old in list(borders):
        borders.remove(old)
    for side in ("top", "left", "bottom", "right"):
        _insert_ordered(borders, _border_el(side, size_pt, color))
    for side in ("insideH", "insideV"):
        _insert_ordered(borders, _border_el(side, size_pt if inside else None, color))


def _cell_borders(cell: _Cell, **sides: tuple[float | None, str]) -> None:
    """Cell borders by side (top/left/bottom/right/...), written in OOXML schema order."""
    borders = OxmlElement("w:tcBorders")
    for side, (size_pt, color) in sides.items():
        _insert_ordered(borders, _border_el(side, size_pt, color))
    _insert_ordered(cell._tc.get_or_add_tcPr(), borders)


def _cell_margins(table: Table, top: int = 40, bottom: int = 40, left: int = 100, right: int = 100) -> None:
    """Default cell padding in twentieths of a point (dxa)."""
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    _insert_ordered(table._tbl.tblPr, mar)


def _fix_table_layout(table: Table, widths: list[Length]) -> None:
    """Fixed layout with explicit column widths so Word and LibreOffice agree."""
    table.autofit = False
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = _insert_ordered(tblpr, OxmlElement("w:tblW"))
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(int(sum(w.twips for w in widths))))
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[min(idx, len(widths) - 1)]
    for col, w in zip(table.columns, widths):
        col.width = w


def _row_no_split(row: Any, header: bool = False) -> None:
    trpr = row._tr.get_or_add_trPr()
    _insert_ordered(trpr, OxmlElement("w:cantSplit"))
    if header:
        _insert_ordered(trpr, OxmlElement("w:tblHeader"))


def _add_field(paragraph: Paragraph, instr: str, **style: Any) -> None:
    """Append a complex field (e.g. PAGE, NUMPAGES) as begin/instr/separate/result/end runs."""

    def _fld_char(kind: str) -> Any:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el

    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    parts = (_fld_char("begin"), instr_el, _fld_char("separate"), None, _fld_char("end"))
    for part in parts:
        run = _style_run(paragraph.add_run("1" if part is None else ""), **style)
        if part is not None:
            run._r.append(part)


def _new_list_num_id(doc: Any) -> int | None:
    """Create a fresh w:num (restarting at 1) based on the 'List Number' style's list."""
    style = doc.styles["List Number"]
    numpr = style.element.pPr.numPr if style.element.pPr is not None else None
    if numpr is None or numpr.numId is None:
        return None
    numbering = doc.part.numbering_part.element
    num = numbering.num_having_numId(numpr.numId.val)
    new_num = numbering.add_num(num.abstractNumId.val)
    new_num.add_lvlOverride(ilvl=0).add_startOverride(1)
    return int(new_num.numId)


def _align_list_level(doc: Any, style: Any, left: Length, hanging: Length) -> None:
    """Copy a list style's indents into its numbering level (ilvl 0).

    Word lets the numbering level's indent and ``num`` tab override the
    paragraph style, LibreOffice does so only when ``numPr`` is set on the
    paragraph itself: with the same values in both places the marker, the
    first line and the wrapped lines line up in both applications.
    """
    ppr_style = style.element.pPr
    numpr = ppr_style.numPr if ppr_style is not None else None
    if numpr is None or numpr.numId is None:
        return
    numbering = doc.part.numbering_part.element
    abstract_id = str(numbering.num_having_numId(numpr.numId.val).abstractNumId.val)
    for abstract in numbering.findall(qn("w:abstractNum")):
        if abstract.get(qn("w:abstractNumId")) != abstract_id:
            continue
        for lvl in abstract.findall(qn("w:lvl")):
            if lvl.get(qn("w:ilvl")) != "0":
                continue
            ppr = lvl.find(qn("w:pPr"))
            if ppr is None:
                ppr = OxmlElement("w:pPr")
                lvl.append(ppr)
            for tabs in ppr.findall(qn("w:tabs")):
                for tab in tabs.findall(qn("w:tab")):
                    tab.set(qn("w:pos"), str(left.twips))
            ind = ppr.find(qn("w:ind"))
            if ind is None:
                ind = _insert_ordered(ppr, OxmlElement("w:ind"))
            ind.attrib.pop(qn("w:firstLine"), None)
            ind.set(qn("w:left"), str(left.twips))
            ind.set(qn("w:hanging"), str(hanging.twips))


def _set_list_num_id(paragraph: Paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    p_numpr = ppr.get_or_add_numPr()
    p_numpr.get_or_add_ilvl().val = 0
    p_numpr.get_or_add_numId().val = num_id


def _spacing(paragraph: Paragraph, before: float | None = None, after: float | None = None, line: float | None = None) -> None:
    pf = paragraph.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line


# ----------------------------------------------------------------------------
# the renderer
# ----------------------------------------------------------------------------


class _Renderer:
    """Builds one branded DOCX. Not reusable."""

    def __init__(self, document: Document, brand: Brand) -> None:
        self.d = document
        self.b = brand
        self.labels = brand.labels_for(document.language)
        self.doc = DocxDocument()
        self.body_font = brand.fonts.body
        self.head_font = brand.fonts.heading
        self.cjk = brand.fonts.body_cjk
        self.c = brand.colors
        page_w, _ = _PAGE_SIZES_CM.get(brand.page.size.upper(), _PAGE_SIZES_CM["A4"])
        m = brand.page.margins_cm
        self.text_width: Length = Cm(page_w - m.left - m.right)

    # -- orchestration -------------------------------------------------------

    def build(self) -> None:
        self._setup_page()
        self._setup_styles()
        self._build_header()
        self._build_footer()
        self._build_cover()
        if self.d.intro:
            self._paragraph(self.d.intro, before=4)
        for section in self.d.sections:
            self._build_section(section)
        if self.d.closing:
            closing = self._paragraph(self.d.closing, before=8)
            closing.paragraph_format.keep_with_next = bool(self.d.signatures)
        if self.d.signatures:
            self._build_signatures()
        if self.d.attachments:
            self._build_attachments()
        if self.d.notes:
            self._build_notes()
        self._set_core_properties()

    def save(self, out_path: Path) -> None:
        buf = io.BytesIO()
        self.doc.save(buf)
        _repack_deterministic(buf.getvalue(), out_path)

    # -- page & styles -------------------------------------------------------

    def _setup_page(self) -> None:
        sec = self.doc.sections[0]
        w, h = _PAGE_SIZES_CM.get(self.b.page.size.upper(), _PAGE_SIZES_CM["A4"])
        sec.page_width, sec.page_height = Cm(w), Cm(h)
        m = self.b.page.margins_cm
        sec.top_margin, sec.bottom_margin = Cm(m.top), Cm(m.bottom)
        sec.left_margin, sec.right_margin = Cm(m.left), Cm(m.right)
        sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.9)

    def _setup_styles(self) -> None:
        styles = self.doc.styles
        lang = {"it": "it-IT", "en": "en-US", "zh": "zh-CN"}.get(self.d.language, "it-IT")

        normal = styles["Normal"]
        self._style_font(normal, self.body_font, 11, self.c.text)
        _set_lang(normal.element, lang)
        _spacing(normal, after=6, line=1.15)

        for name, size, before, after in (("Heading 1", 14, 16, 6), ("Heading 2", 12, 12, 4), ("Heading 3", 11, 10, 3)):
            st = styles[name]
            self._style_font(st, self.head_font, size, self.c.primary, bold=True)
            _spacing(st, before=before, after=after, line=1.1)
            st.paragraph_format.keep_with_next = True

        for name in ("List Bullet", "List Number"):
            st = styles[name]
            self._style_font(st, self.body_font, 11, self.c.text)
            _spacing(st, after=3, line=1.15)
            st.paragraph_format.left_indent = Cm(_LIST_LEFT_CM)
            st.paragraph_format.first_line_indent = Cm(-_LIST_HANGING_CM)
            _align_list_level(self.doc, st, Cm(_LIST_LEFT_CM), Cm(_LIST_HANGING_CM))

        for name in ("Header", "Footer"):
            st = styles[name]
            self._style_font(st, self.body_font, 8, self.c.muted)
            _spacing(st, after=0, line=1.0)

    def _style_font(self, style: Any, font: str, size: float, color: str, bold: bool | None = None) -> None:
        _set_fonts(style.element, font, self.cjk)
        style.font.size = Pt(size)
        style.font.color.rgb = _rgb(color)
        if bold is not None:
            style.font.bold = bold
        style.font.italic = False

    # -- header / footer -----------------------------------------------------

    def _build_header(self) -> None:
        header = self.doc.sections[0].header
        header.is_linked_to_previous = False
        rule = header.paragraphs[0]
        rule.style = self.doc.styles["Header"]
        table = header.add_table(rows=1, cols=2, width=self.text_width)
        rule._p.addprevious(table._tbl)
        widths = [Cm(self.text_width.cm * 0.55), Cm(self.text_width.cm * 0.45)]
        _fix_table_layout(table, widths)
        _table_borders(table, None, self.c.muted)
        _cell_margins(table, 0, 0, 0, 0)
        left, right = table.rows[0].cells
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        self._header_logo(left.paragraphs[0])
        self._header_company(right)
        _spacing(rule, before=2, after=8, line=1.0)
        _style_run(rule.add_run(""), size=2)
        _para_border(rule, "bottom", 1.5, self.c.primary, space=0)

    def _header_logo(self, paragraph: Paragraph) -> None:
        paragraph.style = self.doc.styles["Header"]
        _spacing(paragraph, before=0, after=0, line=1.0)
        logo = self.b.logo_path
        if logo.exists():
            try:
                width, height = self._logo_size(logo)
                paragraph.add_run().add_picture(str(logo), width=width, height=height)
                return
            except Exception as exc:  # unreadable image: fall back to the wordmark
                log.warning("Logo %s non utilizzabile (%s): uso il nome azienda", logo, exc)
        _style_run(paragraph.add_run(self.b.company.name), font=self.head_font, size=16, color=self.c.primary, bold=True, cjk=self.cjk)

    def _logo_size(self, logo: Path) -> tuple[Length, Length]:
        """Logo box: ``width_cm`` wide, shrunk (aspect ratio kept) when taller than ``max_height_cm``.

        Without the cap a square or portrait logo would make the header, which
        repeats on every page, several centimetres tall.
        """
        from docx.image.image import Image as DocxImage

        img = DocxImage.from_file(str(logo))
        ratio = img.px_height / img.px_width
        width = float(self.b.logo.width_cm)
        height = width * ratio
        max_h = float(self.b.logo.max_height_cm)
        if max_h > 0 and height > max_h:
            width, height = max_h / ratio, max_h
        return Cm(width), Cm(height)

    def _header_company(self, cell: _Cell) -> None:
        p = cell.paragraphs[0]
        p.style = self.doc.styles["Header"]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _spacing(p, before=0, after=0, line=1.1)
        _style_run(p.add_run(self.b.company.name), font=self.head_font, size=9, color=self.c.primary, bold=True, cjk=self.cjk)
        if self.b.company.tagline:
            r = p.add_run()
            r.add_break()
            _style_run(p.add_run(self.b.company.tagline), font=self.body_font, size=9, color=self.c.muted, cjk=self.cjk)
        if self.d.confidential and self.b.footer.confidential_label:
            r = p.add_run()
            r.add_break()
            _style_run(p.add_run(self.b.footer.confidential_label.upper()), font=self.head_font, size=7.5, color=self.c.secondary, bold=True, cjk=self.cjk)

    def _build_footer(self) -> None:
        footer = self.doc.sections[0].footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.style = self.doc.styles["Footer"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(p, before=0, after=2, line=1.0)
        _para_border(p, "top", 0.5, self.c.muted, space=4)
        _style_run(p.add_run(self.b.footer.text or self.b.company.name), size=8, color=self.c.muted, cjk=self.cjk)

        show_conf = self.d.confidential and bool(self.b.footer.confidential_label)
        if not (show_conf or self.b.footer.page_numbers):
            return
        # plain paragraph: the built-in Footer style carries its own tab stops
        p2 = footer.add_paragraph()
        _spacing(p2, before=0, after=0, line=1.0)
        p2.paragraph_format.tab_stops.add_tab_stop(self.text_width, WD_TAB_ALIGNMENT.RIGHT)
        if show_conf:
            _style_run(p2.add_run(self.b.footer.confidential_label), size=8, color=self.c.secondary, bold=True, cjk=self.cjk)
        if self.b.footer.page_numbers:
            p2.add_run("\t")
            self._page_numbers(p2)

    def _page_numbers(self, paragraph: Paragraph) -> None:
        lab = self.labels
        zh = self.d.language == "zh"
        style = {"size": 8, "color": self.c.muted, "cjk": self.cjk}
        _style_run(paragraph.add_run(lab["page"] + ("" if zh else " ")), **style)
        _add_field(paragraph, "PAGE", **style)
        _style_run(paragraph.add_run(("" if zh else " ") + lab["of"] + ("" if zh else " ")), **style)
        _add_field(paragraph, "NUMPAGES", **style)
        if zh:
            _style_run(paragraph.add_run("页"), **style)

    # -- cover block ---------------------------------------------------------

    def _build_cover(self) -> None:
        d = self.d
        title = self.doc.add_paragraph()
        _spacing(title, before=4, after=2, line=1.05)
        title.paragraph_format.keep_with_next = True
        add_rich_text(title, d.title, font=self.head_font, size=22, color=self.c.primary, bold=True, cjk=self.cjk)
        if d.subtitle:
            sub = self.doc.add_paragraph()
            _spacing(sub, before=0, after=2, line=1.1)
            sub.paragraph_format.keep_with_next = True
            add_rich_text(sub, d.subtitle, font=self.head_font, size=13, color=self.c.muted, cjk=self.cjk)
        accent = self.doc.add_paragraph()
        _spacing(accent, before=0, after=10, line=1.0)
        accent.paragraph_format.right_indent = Cm(self.text_width.cm - 4.0)
        _style_run(accent.add_run(""), size=2)
        _para_border(accent, "bottom", 2.0, self.c.secondary, space=0)

        if d.recipient and d.recipient.lines():
            table = self._plain_table(1, 2, [0.5, 0.5])
            info_cell, rec_cell = table.rows[0].cells
            self._info_lines(info_cell, first=info_cell.paragraphs[0])
            self._recipient_block(rec_cell)
            self._spacer(6)
        else:
            self._info_lines(None, first=None)

        if d.parties:
            self._parties_table()
        if d.key_facts:
            self._key_facts_table()

    def _info_lines(self, cell: _Cell | None, first: Paragraph | None) -> None:
        d, lab = self.d, self.labels
        rows: list[tuple[str, str]] = []
        if d.reference:
            rows.append((lab["reference"], d.reference))
        date_txt = format_date(d.date, d.language)
        if d.place:
            rows.append((lab["date"], f"{d.place}, {date_txt}"))
        else:
            rows.append((lab["date"], date_txt))
        muted = {"size": 10, "color": self.c.muted, "cjk": self.cjk}
        normal = {"size": 10, "color": self.c.text, "cjk": self.cjk}
        for i, (label, value) in enumerate(rows):
            p = first if (i == 0 and first is not None) else (cell.add_paragraph() if cell else self.doc.add_paragraph())
            _spacing(p, before=0, after=1, line=1.1)
            _style_run(p.add_run(label + "  "), **muted)
            _style_run(p.add_run(value), **normal)
        if d.status:
            p = cell.add_paragraph() if cell else self.doc.add_paragraph()
            _spacing(p, before=4, after=8 if cell is None else 0, line=1.1)
            self._status_tag(p)
        elif cell is None:
            _spacing(p, after=10)

    def _status_tag(self, paragraph: Paragraph) -> None:
        status = (self.d.status or "").strip()
        is_draft = status.lower() in {"bozza", "draft", "草稿", self.labels["draft"].lower()}
        fill = self.c.secondary if is_draft else _blend(self.c.primary, 0.15)
        color = "#FFFFFF" if is_draft else self.c.primary
        run = paragraph.add_run(f"  {status.upper()}  ")
        _style_run(run, font=self.head_font, size=8, color=color, bold=True, cjk=self.cjk)
        _run_shade(run, fill)

    def _recipient_block(self, cell: _Cell) -> None:
        rec = self.d.recipient
        assert rec is not None
        p = cell.paragraphs[0]
        _spacing(p, before=0, after=1, line=1.1)
        _style_run(p.add_run(self.labels["recipient"]), size=9, color=self.c.muted, cjk=self.cjk)
        for i, line in enumerate(rec.lines()):
            q = cell.add_paragraph()
            _spacing(q, before=0, after=0, line=1.15)
            bold = i == 0
            add_rich_text(q, line, size=10.5, color=self.c.text, bold=bold, cjk=self.cjk)

    def _parties_table(self) -> None:
        table = self._plain_table(len(self.d.parties), 2, [0.24, 0.76])
        _table_borders(table, 0.5, _blend(self.c.muted, 0.35), inside=True)
        _cell_margins(table, 50, 50, 110, 110)
        for row, party in zip(table.rows, self.d.parties):
            _row_no_split(row)
            role_cell, name_cell = row.cells
            _shade(role_cell, _blend(self.c.primary, 0.12))
            role_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = role_cell.paragraphs[0]
            _spacing(p, after=0, line=1.1)
            _style_run(p.add_run(party.role), font=self.head_font, size=9.5, color=self.c.primary, bold=True, cjk=self.cjk)
            p = name_cell.paragraphs[0]
            _spacing(p, after=0, line=1.1)
            add_rich_text(p, party.name, size=10.5, color=self.c.text, bold=True, cjk=self.cjk)
            if party.details:
                q = name_cell.add_paragraph()
                _spacing(q, after=0, line=1.1)
                add_rich_text(q, party.details, size=9, color=self.c.muted, cjk=self.cjk)
        self._spacer(6)

    def _key_facts_table(self) -> None:
        table = self._plain_table(len(self.d.key_facts), 2, [0.30, 0.70])
        _table_borders(table, 0.5, _blend(self.c.muted, 0.35), inside=True)
        _cell_margins(table, 45, 45, 110, 110)
        for row, (label, value) in zip(table.rows, self.d.key_facts):
            _row_no_split(row)
            lab_cell, val_cell = row.cells
            _shade(lab_cell, _blend(self.c.primary, 0.20))
            p = lab_cell.paragraphs[0]
            _spacing(p, after=0, line=1.1)
            add_rich_text(p, label, size=10, color=self.c.primary, bold=True, cjk=self.cjk)
            p = val_cell.paragraphs[0]
            _spacing(p, after=0, line=1.1)
            add_rich_text(p, value, size=10, color=self.c.text, cjk=self.cjk)
        self._spacer(8)

    # -- body ----------------------------------------------------------------

    def _build_section(self, section: Section) -> None:
        style = "Heading 1" if section.level == 1 else "Heading 2"
        h = self.doc.add_paragraph(style=style)
        add_rich_text(h, section.heading, cjk=self.cjk)
        for block in section.blocks:
            self._build_block(block)

    def _build_block(self, block: Block) -> None:
        kind = block.type
        if kind == "paragraph":
            self._paragraph(block.text or "")
        elif kind == "bullets":
            self._list(block.items, "List Bullet")
        elif kind == "numbered":
            self._list(block.items, "List Number")
        elif kind == "table":
            self._data_table(block)
        elif kind == "callout":
            self._callout(block)
        elif kind == "quote":
            self._quote(block.text or "")
        elif kind == "page_break":
            self.doc.add_page_break()

    def _paragraph(self, text: str, before: float | None = None, after: float | None = None) -> Paragraph:
        p = self.doc.add_paragraph()
        # justified text stretches lines that end with a manual break: keep those left-aligned
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if "\n" in text else WD_ALIGN_PARAGRAPH.JUSTIFY
        _spacing(p, before=before, after=after)
        add_rich_text(p, text, cjk=self.cjk)
        return p

    def _list(self, items: Iterable[str], style: str) -> None:
        num_id = _new_list_num_id(self.doc) if style == "List Number" else None
        last: Paragraph | None = None
        for item in items:
            last = self.doc.add_paragraph(style=style)
            if num_id is not None:
                _set_list_num_id(last, num_id)
            add_rich_text(last, item, cjk=self.cjk)
        if last is not None:
            _spacing(last, after=6)

    def _data_table(self, block: Block) -> None:
        ncols = max(1, len(block.columns))
        widths = self._col_fractions(block.col_widths, ncols)
        table = self._plain_table(1 + len(block.rows), ncols, widths)
        _table_borders(table, 0.5, _blend(self.c.muted, 0.45), inside=True)
        _cell_margins(table, 45, 45, 90, 90)
        head = table.rows[0]
        _row_no_split(head, header=True)
        for cell, name in zip(head.cells, block.columns):
            _shade(cell, self.c.primary)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            _spacing(p, after=0, line=1.1)
            add_rich_text(p, name, size=9.5, color=self.c.table_header_text, bold=True, cjk=self.cjk)
        for r_idx, (row, values) in enumerate(zip(table.rows[1:], block.rows)):
            _row_no_split(row)
            for c_idx, cell in enumerate(row.cells):
                if r_idx % 2 == 1:
                    _shade(cell, self.c.table_alt_row)
                value = values[c_idx] if c_idx < len(values) else ""
                p = cell.paragraphs[0]
                _spacing(p, after=0, line=1.1)
                add_rich_text(p, value, size=9.5, color=self.c.text, cjk=self.cjk)
        if block.caption:
            cap = self.doc.add_paragraph()
            _spacing(cap, before=3, after=8, line=1.1)
            add_rich_text(cap, block.caption, size=8.5, color=self.c.muted, italic=True, cjk=self.cjk)
        else:
            self._spacer(8)

    def _callout(self, block: Block) -> None:
        style = block.style or "info"
        if style == "warning":
            fill, edge = _blend(self.c.secondary, 0.14), self.c.secondary
        elif style == "note":
            fill, edge = self.c.table_alt_row, self.c.muted
        else:
            fill, edge = self.c.callout_bg, self.c.secondary
        table = self._plain_table(1, 1, [1.0])
        _cell_margins(table, 90, 90, 170, 140)
        cell = table.rows[0].cells[0]
        _row_no_split(table.rows[0])
        _shade(cell, fill)
        _cell_borders(cell, left=(3.0, edge), top=(None, edge), bottom=(None, edge), right=(None, edge))
        p = cell.paragraphs[0]
        _spacing(p, after=0, line=1.15)
        add_rich_text(p, block.text or "", size=10.5, color=self.c.text, cjk=self.cjk)
        self._spacer(8)

    def _quote(self, text: str) -> None:
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent, pf.right_indent = Cm(1.0), Cm(0.8)
        _spacing(p, before=4, after=8, line=1.15)
        _para_border(p, "left", 1.5, self.c.secondary, space=8)
        add_rich_text(p, text, size=10.5, color=self.c.muted, italic=True, cjk=self.cjk)

    # -- tail ----------------------------------------------------------------

    def _build_signatures(self) -> None:
        h = self.doc.add_paragraph(style="Heading 2")
        _spacing(h, before=18)
        add_rich_text(h, self.labels["signatures"], cjk=self.cjk)
        sigs = list(self.d.signatures)
        per_row = 3 if len(sigs) > 2 else 2  # a lone signatory keeps a half-width line
        for start in range(0, len(sigs), per_row):
            chunk = sigs[start : start + per_row]
            table = self._plain_table(1, per_row, [1.0 / per_row] * per_row)
            _cell_margins(table, 0, 0, 0, 170)
            _row_no_split(table.rows[0])
            for cell, sig in zip(table.rows[0].cells, chunk):
                self._signature_cell(cell, sig)
            self._spacer(6)

    def _signature_cell(self, cell: _Cell, sig: Any) -> None:
        p = cell.paragraphs[0]
        _spacing(p, before=6, after=0, line=1.1)
        add_rich_text(p, sig.party, size=10, color=self.c.primary, bold=True, cjk=self.cjk)
        line = cell.add_paragraph()
        _spacing(line, before=30, after=2, line=1.0)
        line.paragraph_format.right_indent = Cm(0.4)
        _style_run(line.add_run(""), size=2)
        _para_border(line, "bottom", 0.75, self.c.text, space=1)
        name = cell.add_paragraph()
        _spacing(name, before=0, after=0, line=1.1)
        add_rich_text(name, sig.name or " ", size=10, color=self.c.text, cjk=self.cjk)
        if sig.title:
            title = cell.add_paragraph()
            _spacing(title, before=0, after=0, line=1.1)
            add_rich_text(title, sig.title, size=9, color=self.c.muted, cjk=self.cjk)

    def _build_attachments(self) -> None:
        h = self.doc.add_paragraph(style="Heading 2")
        add_rich_text(h, self.labels["attachments"], cjk=self.cjk)
        self._list(self.d.attachments, "List Bullet")

    def _build_notes(self) -> None:
        p = self.doc.add_paragraph()
        _spacing(p, before=14, after=0, line=1.1)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _para_border(p, "top", 0.5, _blend(self.c.muted, 0.5), space=6)
        add_rich_text(p, self.d.notes or "", size=8.5, color=self.c.muted, cjk=self.cjk)

    def _set_core_properties(self) -> None:
        cp = self.doc.core_properties
        d = self.d
        cp.title = d.title
        cp.author = self.b.company.name
        cp.last_modified_by = self.b.company.name
        cp.subject = d.doc_type
        cp.keywords = d.document_id
        entry = DOC_TYPES.get(d.doc_type)
        cp.category = entry.name_it if entry else d.doc_type
        cp.comments = d.subtitle or ""
        cp.language = {"it": "it-IT", "en": "en-US", "zh": "zh-CN"}.get(d.language, "it-IT")
        stamp = dt.datetime.combine(d.date_value, dt.time(0, 0))
        cp.created = stamp
        cp.modified = stamp
        cp.revision = 1

    # -- shared building blocks ---------------------------------------------

    def _plain_table(self, rows: int, cols: int, fractions: list[float]) -> Table:
        """Borderless, fixed-width table spanning the text width."""
        table = self.doc.add_table(rows=rows, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        total = self.text_width.cm
        widths = [Cm(total * f) for f in fractions]
        _fix_table_layout(table, widths)
        _table_borders(table, None, self.c.muted)
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _spacing(p, before=0, after=0, line=1.1)
        return table

    def _col_fractions(self, col_widths: list[float] | None, ncols: int) -> list[float]:
        if col_widths and len(col_widths) == ncols and sum(col_widths) > 0:
            total = float(sum(col_widths))
            return [w / total for w in col_widths]
        return [1.0 / ncols] * ncols

    def _spacer(self, after: float) -> Paragraph:
        """Tiny paragraph used as vertical space after tables."""
        p = self.doc.add_paragraph()
        _spacing(p, before=0, after=after, line=1.0)
        _style_run(p.add_run(""), size=4)
        return p


# ----------------------------------------------------------------------------
# deterministic zip
# ----------------------------------------------------------------------------


def _repack_deterministic(data: bytes, out_path: Path) -> None:
    """Rewrite the DOCX zip with fixed timestamps so identical input → identical bytes."""
    src = zipfile.ZipFile(io.BytesIO(data))
    with zipfile.ZipFile(out_path, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=_FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = item.external_attr
            dst.writestr(info, src.read(item.filename))
