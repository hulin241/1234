"""Plain-text extraction from office files sent by users.

Supported: .docx (paragraphs + tables as pipe rows), .xlsx (one block per
sheet, rows as TSV), .csv/.txt/.md (read as text).  Anything else, including
.pptx, raises :class:`UnsupportedFileError`.  Output longer than ``max_chars``
is cut and ends with ``TRUNCATION_MARKER``.
"""

from __future__ import annotations

from pathlib import Path

TRUNCATION_MARKER = "[… troncato]"
MAX_SHEET_ROWS = 500
_TEXT_ENCODINGS = ("utf-8-sig", "cp1252")  # tried strictly, in order


class UnsupportedFileError(ValueError):
    """The file type cannot be read by ``extract_text``."""


def extract_text(path: Path | str, max_chars: int = 60_000) -> str:
    """Return the text content of ``path`` (see module docstring)."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"File non trovato: {path}")
    suffix = path.suffix.lower()
    if suffix == ".docx":
        text = _docx_text(path)
    elif suffix == ".xlsx":
        text = _xlsx_text(path)
    elif suffix in (".csv", ".txt", ".md"):
        text = _plain_text(path)
    else:
        raise UnsupportedFileError(
            f"Tipo di file non supportato: '{suffix or path.name}' (leggo .docx, .xlsx, .csv, .txt, .md)"
        )
    return _truncate(text, max_chars)


def _docx_text(path: Path) -> str:
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = DocxDocument(str(path))
    lines: list[str] = []
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            text = Paragraph(el, doc).text.strip()
            if text:
                lines.append(text)
        elif el.tag == qn("w:tbl"):
            for row in Table(el, doc).rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                lines.append("| " + " | ".join(cells) + " |")
            lines.append("")
    return "\n".join(lines).strip()


def _xlsx_text(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[str] = []
    try:
        for ws in wb.worksheets:
            lines = [f"## Foglio: {ws.title}"]
            count = 0
            for row in ws.iter_rows(values_only=True):
                if all(v is None or str(v).strip() == "" for v in row):
                    continue
                count += 1
                if count > MAX_SHEET_ROWS:
                    lines.append(f"[… altre righe omesse oltre le prime {MAX_SHEET_ROWS}]")
                    break
                lines.append("\t".join("" if v is None else str(v) for v in row))
            blocks.append("\n".join(lines))
    finally:
        wb.close()
    return "\n\n".join(blocks).strip()


def _plain_text(path: Path) -> str:
    """Decode UTF-8 (BOM tolerated) or, failing that, cp1252 (Excel's plain CSV on Windows)."""
    raw = path.read_bytes()
    for encoding in _TEXT_ENCODINGS:
        try:
            return raw.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace").strip()


def _truncate(text: str, max_chars: int) -> str:
    if max_chars <= 0 or len(text) <= max_chars:
        return text
    keep = max(0, max_chars - len(TRUNCATION_MARKER) - 1)
    return text[:keep].rstrip() + "\n" + TRUNCATION_MARKER
